"""
Gera novas planilhas XLSX usando o modelo original como template.
Preserva estilos, mesclagens, bordas, fontes e cores da planilha modelo.
"""
from __future__ import annotations

import copy
from datetime import date
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from loguru import logger

import pandas as pd


# ---------------------------------------------------------------------------
# Helpers de cópia de estilo
# ---------------------------------------------------------------------------

def _copy_cell_style(src_cell, dst_cell) -> None:
    """Copia todos os atributos de estilo de uma célula para outra."""
    if src_cell.has_style:
        dst_cell.font      = copy.copy(src_cell.font)
        dst_cell.border    = copy.copy(src_cell.border)
        dst_cell.fill      = copy.copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy.copy(src_cell.protection)
        dst_cell.alignment  = copy.copy(src_cell.alignment)


def _copy_sheet_structure(src_ws, dst_ws) -> None:
    """Copia estrutura visual completa: células, estilos, mesclagens, larguras."""
    # Largura de colunas
    for col_letter, col_dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col_letter].width = col_dim.width

    # Altura de linhas
    for row_num, row_dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[row_num].height = row_dim.height

    # Células
    for row in src_ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            dst_cell = dst_ws.cell(row=cell.row, column=cell.column)
            dst_cell.value = cell.value
            _copy_cell_style(cell, dst_cell)

    # Mesclagens
    for merge_range in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merge_range))

    # Zoom e congelamento
    dst_ws.sheet_view.zoomScale = src_ws.sheet_view.zoomScale
    if src_ws.freeze_panes:
        dst_ws.freeze_panes = src_ws.freeze_panes


# ---------------------------------------------------------------------------
# Status → cor de preenchimento
# ---------------------------------------------------------------------------
_STATUS_COLORS = {
    "verde":    "92D050",  # verde
    "amarelo":  "FFFF00",  # amarelo
    "vermelho": "FF0000",  # vermelho
}


def _status_fill(status: Optional[str]) -> Optional[PatternFill]:
    if not status:
        return None
    cor = _STATUS_COLORS.get(str(status).lower())
    if cor:
        return PatternFill("solid", fgColor=cor)
    return None


# ---------------------------------------------------------------------------
# Exportador principal
# ---------------------------------------------------------------------------

class Exportador:
    def __init__(self, template_path: str = "templates/modelo.xlsx"):
        self.template_path = Path(template_path)

    def _load_template(self) -> openpyxl.Workbook:
        if not self.template_path.exists():
            logger.warning(f"Template não encontrado: {self.template_path}. Usando workbook em branco.")
            return openpyxl.Workbook()
        return load_workbook(str(self.template_path))

    def gerar_matriz(
        self,
        df: pd.DataFrame,
        nome_matriz: str,
        tipo_quadro: str,
        data_referencia: date,
        output_path: Optional[str] = None,
    ) -> Path:
        """
        Gera um arquivo XLSX para uma matriz específica.
        Usa o template como base visual e preenche com os dados do DataFrame.
        """
        wb_template = self._load_template()

        wb_out = openpyxl.Workbook()
        wb_out.remove(wb_out.active)  # remove aba padrão

        sheet_name = f"{tipo_quadro} - {nome_matriz} {data_referencia.strftime('%d.%m')}"[:31]

        # Se o template tem uma aba com estrutura similar, usa ela; senão cria nova
        src_ws_name = None
        for name in wb_template.sheetnames:
            if "quadro" in name.lower():
                src_ws_name = name
                break

        if src_ws_name:
            src_ws = wb_template[src_ws_name]
            dst_ws = wb_out.create_sheet(sheet_name)
            _copy_sheet_structure(src_ws, dst_ws)
        else:
            dst_ws = wb_out.create_sheet(sheet_name)
            self._create_blank_sheet(dst_ws, nome_matriz, tipo_quadro, data_referencia)

        # Preenche dados (tenta localizar header automaticamente)
        self._fill_data(dst_ws, df)

        if not output_path:
            Path("outputs").mkdir(parents=True, exist_ok=True)
            fname = f"matriz_{nome_matriz}_{tipo_quadro}_{data_referencia.strftime('%Y%m%d')}.xlsx"
            output_path = Path("outputs") / fname

        wb_out.save(str(output_path))
        logger.info(f"Exportado: {output_path}")
        return Path(output_path)

    def _create_blank_sheet(self, ws, nome_matriz: str, tipo_quadro: str, data_referencia: date):
        """Cria aba com layout padrão quando não há template."""
        header_fill = PatternFill("solid", fgColor="1F3864")
        header_font = Font(bold=True, color="FFFFFF", size=11)

        # Título
        ws.merge_cells("A1:I1")
        title = ws["A1"]
        title.value = f"{tipo_quadro} — {nome_matriz} — {data_referencia.strftime('%d/%m/%Y')}"
        title.font = Font(bold=True, size=14, color="1F3864")
        title.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 30

        # Cabeçalhos
        headers = ["Categoria", "Indicador", "Meta", "Realizado", "Resultado %",
                   "Status", "Área", "Setor", "Responsável", "Observação"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = 18

        ws.freeze_panes = "A4"

    def _fill_data(self, ws, df: pd.DataFrame):
        """
        Localiza o cabeçalho na planilha e escreve os dados abaixo dele.
        Se não encontrar, escreve a partir da linha 4.
        """
        # Localiza linha de cabeçalho com a coluna "indicador"
        header_row = None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and "indicador" in str(cell.value).lower():
                    header_row = cell.row
                    break
            if header_row:
                break

        start_row = (header_row + 1) if header_row else 4

        # Mapa coluna_nome → índice (procura no header_row)
        col_map = {}
        if header_row:
            for cell in ws[header_row]:
                if cell.value:
                    col_map[str(cell.value).lower().strip()] = cell.column

        col_fields = {
            "categoria":  next((col_map[k] for k in col_map if "categ" in k), 1),
            "indicador":  next((col_map[k] for k in col_map if "indic" in k), 2),
            "meta":       next((col_map[k] for k in col_map if "meta" in k), 3),
            "realizado":  next((col_map[k] for k in col_map if "realiz" in k or "result" in k), 4),
            "status":     next((col_map[k] for k in col_map if "status" in k or "sem" in k), 5),
            "area":       next((col_map[k] for k in col_map if "area" in k or "área" in k), 6),
            "setor":      next((col_map[k] for k in col_map if "setor" in k), 7),
            "responsavel":next((col_map[k] for k in col_map if "resp" in k), 8),
            "observacao": next((col_map[k] for k in col_map if "obs" in k), 9),
        }

        thin = Side(style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for i, (_, row) in enumerate(df.iterrows()):
            r = start_row + i
            for field, col_idx in col_fields.items():
                cell = ws.cell(row=r, column=col_idx, value=row.get(field))
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

                # Cor de status
                if field == "status":
                    fill = _status_fill(row.get("status"))
                    if fill:
                        cell.fill = fill
                        cell.font = Font(bold=True, color="000000")

            ws.row_dimensions[r].height = 18

    def exportar_todos(self, db, output_folder: str = "outputs") -> list[Path]:
        """Exporta todas as matrizes do banco para arquivos Excel separados."""
        Path(output_folder).mkdir(parents=True, exist_ok=True)
        df_matrizes = db.get_matrizes()
        paths = []

        for _, m in df_matrizes.iterrows():
            df = db.get_indicadores(
                nome_matriz=m["nome_matriz"],
                tipo_quadro=m["tipo_quadro"],
                data_de=str(m["data_referencia"]),
                data_ate=str(m["data_referencia"]),
            )
            if df.empty:
                continue
            try:
                path = self.gerar_matriz(
                    df=df,
                    nome_matriz=m["nome_matriz"],
                    tipo_quadro=m["tipo_quadro"],
                    data_referencia=m["data_referencia"],
                    output_path=Path(output_folder) / f"matriz_{m['nome_matriz']}_{m['data_referencia']}.xlsx",
                )
                paths.append(path)
            except Exception as exc:
                logger.error(f"Erro ao exportar {m['nome_matriz']}: {exc}")

        return paths
